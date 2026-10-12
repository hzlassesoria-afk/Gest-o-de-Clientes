import json, subprocess, sys, datetime, glob
from openpyxl import load_workbook
SK = glob.glob('/root/.claude/skills/synced/*/xlsx')[0]
MODE = {'Custo': 'custo', 'Porcentagem': 'taxa', 'Número': 'numero'}

def run(name, month, inv, stages, comp, today, entries):
    wb = load_workbook('modelo.xlsx')
    b, d = wb['Base de Dados'], wb['Meta Diária']
    b['B4'] = datetime.date(int(month[:4]), int(month[5:]), 1); b['B5'] = inv
    b['B6'] = datetime.date.fromisoformat(today); b['B7'] = 'Sim' if comp else 'Não'
    for i, (modo, val) in enumerate(stages): b[f'B{10+i}'] = modo; b[f'C{10+i}'] = val
    cols = ['G', 'I', 'K', 'M', 'O']
    for day, vals in entries.items():
        for c, v in zip(cols, vals): d[f'{c}{6+day-1}'] = v
    wb.save(f'{name}.xlsx')
    subprocess.run(['python3', SK + '/scripts/recalc.py', f'{name}.xlsx', '120'], capture_output=True, check=True)
    days = {}
    for day, v in entries.items():
        days[f'{month}-{day:02d}'] = {'investimento': v[0], 'leads': v[1], 'cotacoes': v[2], 'negociacoes': v[3], 'vendas': v[4]}
    itens = {'leads': {'modo': MODE[stages[0][0]], 'valor': stages[0][1]}, 'leadsQualificados': {'modo': 'numero', 'valor': None}}
    for k, (modo, val) in zip(['cotacoes', 'negociacoes', 'vendas'], stages[1:]):
        itens[k] = {'modo': MODE[modo], 'valor': val}
    js = f"""
const G=require('/home/user/Gest-o-de-Clientes/public/goals');
const r=G.compute('{month}',{{meta:{{investimento:{inv},itens:{json.dumps(itens)}}},compensarExcedente:{str(comp).lower()},days:{json.dumps(days)}}},'{today}');
const out={{days:{{}},weeks:[]}};
for(const w of r.weeks){{out.weeks.push({{n:w.n,rows:Object.fromEntries(['investimento','leads','cotacoes','negociacoes','vendas'].map(k=>[k,w.rows[k].meta]))}});
 for(const d of w.days) if(d.working) out.days[d.date]=Object.fromEntries(['investimento','leads','cotacoes','negociacoes','vendas'].map(k=>[k,d.rows[k].meta]));}}
console.log(JSON.stringify(out));"""
    eng = json.loads(subprocess.run(['node', '-e', js], capture_output=True, text=True, check=True).stdout)
    wb2 = load_workbook(f'{name}.xlsx', data_only=True)
    d2 = wb2['Meta Diária']; s2 = wb2['Meta Semanal']
    cmap = {'investimento': 'F', 'leads': 'H', 'cotacoes': 'J', 'negociacoes': 'L', 'vendas': 'N'}
    n = diffs = 0
    for r in range(6, 37):
        dt = d2[f'A{r}'].value
        if dt in (None, ''): continue
        key = dt.strftime('%Y-%m-%d')
        if key not in eng['days']:
            if d2[f'E{r}'].value == 'Sim': print('  útil só na planilha', key); diffs += 1
            continue
        for k, c in cmap.items():
            sv = d2[f'{c}{r}'].value; ev = eng['days'][key][k]; n += 1
            if abs((sv or 0) - (ev or 0)) > 1e-6: diffs += 1; print('  DIFF', key, k, sv, ev)
    wcol = {'investimento': 'D', 'leads': 'F', 'cotacoes': 'H', 'negociacoes': 'J', 'vendas': 'L'}
    for w in eng['weeks']:
        for k, c in wcol.items():
            sv = s2[f'{c}{5+w["n"]}'].value; ev = w['rows'][k]; n += 1
            if abs((sv or 0) - (ev or 0)) > 1e-6: diffs += 1; print('  DIFF semana', w['n'], k, sv, ev)
    print(name, 'comparados', n, 'diferenças', diffs)

if __name__ == '__main__':
    base_entries = {1: (1000, 30, 5, 3, 1), 2: (1200, 20, 4, 2, 0), 5: (1500, 35, 6, 4, 1), 6: (900, 25, 3, 2, 0), 7: (2000, 50, 9, 6, 2), 8: (800, 20, 4, 2, 1)}
    st = [('Custo', 35), ('Porcentagem', 20), ('Porcentagem', 60), ('Porcentagem', 30)]
    run('s1', '2026-10', 30000, st, False, '2026-10-14', base_entries)
    run('s2', '2026-10', 30000, st, True, '2026-10-14', base_entries)
    run('s3', '2027-03', 12000, [('Número', 500), ('Custo', 150), ('Número', 40), ('Porcentagem', 25)], False, '2027-03-10',
        {1: (400, 30, 10, 3, 0), 2: (500, 35, 12, 4, 1), 3: (300, 80, 5, 1, 0), 4: (600, 10, 2, 0, 0), 5: (450, 40, 8, 2, 1), 8: (500, 60, 9, 3, 0), 9: (350, 45, 7, 2, 2)})
    run('s4', '2027-03', 12000, [('Número', 500), ('Custo', 150), ('Número', 40), ('Porcentagem', 25)], True, '2027-03-10',
        {1: (400, 30, 10, 3, 0), 2: (500, 35, 12, 4, 1), 3: (300, 80, 5, 1, 0), 4: (600, 10, 2, 0, 0), 5: (450, 40, 8, 2, 1), 8: (500, 60, 9, 3, 0), 9: (350, 45, 7, 2, 2)})
    run('s5', '2026-11', 8000, [('Custo', 40), ('Porcentagem', 15), ('Porcentagem', 50), ('Porcentagem', 20)], False, '2026-11-01', {})
