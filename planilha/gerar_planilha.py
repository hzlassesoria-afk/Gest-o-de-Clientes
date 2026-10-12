"""Gera a planilha Meta Mensal -> Semanal -> Diária (com base de dados de previsões).

Uso: python build_xlsx.py saida.xlsx [--teste]
  --teste: grava lançamentos de exemplo e fixa "hoje" para conferir contra o motor do sistema.
"""
import sys
from datetime import date, timedelta
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter as L
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import FormulaRule
from openpyxl.comments import Comment

OUT = sys.argv[1]
TEST = '--teste' in sys.argv

FONT = 'Arial'
f_norm = Font(name=FONT, size=10)
f_bold = Font(name=FONT, size=10, bold=True)
f_title = Font(name=FONT, size=16, bold=True, color='FF6A1A')
f_sub = Font(name=FONT, size=10, italic=True, color='666666')
f_head = Font(name=FONT, size=10, bold=True, color='FFFFFF')
f_input = Font(name=FONT, size=10, color='0000FF')
f_link = Font(name=FONT, size=10, color='008000')
fill_head = PatternFill('solid', fgColor='0F0E10')
fill_input = PatternFill('solid', fgColor='FFF2A8')
fill_sec = PatternFill('solid', fgColor='FFE3D1')
fill_grey = PatternFill('solid', fgColor='E7E7E7')
thin = Side(style='thin', color='BBBBBB')
box = Border(left=thin, right=thin, top=thin, bottom=thin)
center = Alignment(horizontal='center', vertical='center', wrap_text=True)
left = Alignment(horizontal='left', vertical='center', wrap_text=True)

MONEY = '"R$" #,##0.00'
INT = '#,##0'
PCT = '0.0%'
DATEF = 'dd/mm/yyyy'

wb = Workbook()
base = wb.active
base.title = 'Base de Dados'
mensal = wb.create_sheet('Meta Mensal')
semanal = wb.create_sheet('Meta Semanal')
diario = wb.create_sheet('Meta Diária')
calc = wb.create_sheet('Cálculo')
folgas = wb.create_sheet('Folgas')

B = "'Base de Dados'!"
D = "'Meta Diária'!"
C = "'Cálculo'!"
F = "'Folgas'!"
HOJE = B + '$B$6'
COMP = B + '$B$7="Sim"'
R0, R1 = 6, 36          # linhas dos dias (31 no máximo)


def style_range(ws, rng, font=None, fill=None, fmt=None, align=None, border=None):
    for row in ws[rng]:
        for c in row:
            if font: c.font = font
            if fill: c.fill = fill
            if fmt: c.number_format = fmt
            if align: c.alignment = align
            if border: c.border = border


def head(ws, row, labels, col=1):
    for i, t in enumerate(labels):
        c = ws.cell(row=row, column=col + i, value=t)
        c.font, c.fill, c.alignment, c.border = f_head, fill_head, center, box


def put(ws, ref, value, font=f_norm, fmt=None, fill=None, align=None, border=box):
    c = ws[ref]
    c.value = value
    c.font = font
    if fmt: c.number_format = fmt
    if fill: c.fill = fill
    if align: c.alignment = align
    if border: c.border = border
    return c


# ============================================================ Base de Dados
ws = base
ws['A1'] = 'Base de dados — previsões do mês'; ws['A1'].font = f_title
ws['A2'] = 'Preencha só as células amarelas (letras azuis). Todas as outras abas se calculam sozinhas a partir daqui.'
ws['A2'].font = f_sub

put(ws, 'A4', 'Mês de referência (1º dia do mês)', f_bold)
put(ws, 'B4', date(2026, 10, 1), f_input, 'mm/yyyy', fill_input)
put(ws, 'A5', 'Investimento previsto (R$)', f_bold)
put(ws, 'B5', 30000, f_input, MONEY, fill_input)
put(ws, 'A6', 'Data de hoje (para saber o que já passou)', f_bold)
put(ws, 'B6', date(2026, 10, 14) if TEST else '=TODAY()', f_input if TEST else f_norm, DATEF, fill_input if TEST else None)
put(ws, 'A7', 'Reduzir as metas seguintes quando eu superar a meta?', f_bold)
put(ws, 'B7', 'Não', f_input, None, fill_input, center)
for r, t in [(4, 'Qualquer dia do mês serve; a planilha usa o mês inteiro.'),
             (5, 'Quanto será investido em anúncios no mês.'),
             (6, 'Vem da data do computador. Para simular outro dia, digite a data (azul).'),
             (7, '"Não": a meta só sobe quando fico abaixo. "Sim": ela também cai quando eu supero.')]:
    ws.cell(row=r, column=3, value=t).font = f_sub

head(ws, 9, ['Etapa', 'Como definir', 'Valor previsto', 'Unidade do valor', 'Meta do mês (nº)',
             'Custo por unidade (R$)', '% da etapa anterior', 'Alerta'])
stages = [('Lead', 'Custo', 35), ('Lead em cotação', 'Porcentagem', 20),
          ('Lead em negociação', 'Porcentagem', 60), ('Venda', 'Porcentagem', 30)]
for i, (name, modo, val) in enumerate(stages):
    r = 10 + i
    put(ws, f'A{r}', name, f_bold)
    put(ws, f'B{r}', modo, f_input, None, fill_input, center)
    put(ws, f'C{r}', val, f_input, '#,##0.00', fill_input)
    put(ws, f'D{r}', f'=IF(B{r}="Custo","R$ por "&LOWER(A{r}),IF(B{r}="Porcentagem","% da etapa anterior",IF(B{r}="Número","unidades no mês","")))', f_norm)
    if i == 0:
        e = (f'=IFERROR(IF(C{r}="","",IF(B{r}="Custo",ROUNDUP(ROUND($B$5/C{r},6),0),'
             f'IF(B{r}="Número",ROUNDUP(ROUND(C{r},6),0),""))),"")')
    else:
        e = (f'=IFERROR(IF(C{r}="","",IF(B{r}="Custo",ROUNDUP(ROUND($B$5/C{r},6),0),'
             f'IF(B{r}="Porcentagem",IF(ISNUMBER(E{r-1}),ROUNDUP(ROUND(E{r-1}*C{r}/100,6),0),""),'
             f'IF(B{r}="Número",ROUNDUP(ROUND(C{r},6),0),"")))),"")')
    put(ws, f'E{r}', e, f_norm, INT)
    put(ws, f'F{r}', f'=IF(AND(ISNUMBER(E{r}),$B$5>0),IF(B{r}="Custo",C{r},IF(E{r}>0,$B$5/E{r},"")),"")', f_norm, MONEY)
    if i == 0:
        put(ws, f'G{r}', '—', f_norm, PCT, None, center)
    else:
        put(ws, f'G{r}', f'=IF(AND(ISNUMBER(E{r}),ISNUMBER(E{r-1})),IF(E{r-1}>0,E{r}/E{r-1},""),"")', f_norm, PCT)
    alerta = (f'=IF(C{r}="","Preencha o valor",IF(AND(B{r}="Porcentagem",C{r}>100),"Porcentagem acima de 100",'
              f'IF(AND(B{r}="Porcentagem",NOT(ISNUMBER(E{r-1}))),"Defina a etapa anterior",'
              f'IF(AND(B{r}="Custo",$B$5<=0),"Informe o investimento previsto",IF(E{r}="","Confira o valor","")))))'
              if i else
              f'=IF(C{r}="","Preencha o valor",IF(AND(B{r}="Custo",$B$5<=0),"Informe o investimento previsto",IF(E{r}="","Confira o valor","")))')
    put(ws, f'H{r}', alerta, Font(name=FONT, size=10, color='C00000'))

ws['A15'] = 'Linha de exemplo: Investimento R$ 30.000, Lead a R$ 35,00 (858 leads), cotação em 20% dos leads, negociação em 60% das cotações, venda em 30% das negociações. Troque pelos seus números.'
ws['A15'].font = f_sub
leg = [
    ('Como usar', None),
    ('1.', 'Preencha o mês, o investimento previsto e, para cada etapa, escolha "Custo" (R$ por unidade), "Porcentagem" (% da etapa anterior, digite 20 para 20%) ou "Número" (quantidade exata no mês).'),
    ('2.', 'Lead só aceita Custo ou Número: ele é o início do funil, não tem etapa anterior. As demais etapas aceitam os três.'),
    ('3.', 'Pessoas (leads, cotações, negociações, vendas) são sempre números inteiros: o total do mês é arredondado para cima.'),
    ('4.', 'Em "Meta Diária", digite só o REALIZADO de cada dia (células amarelas). As metas de semana e de dia se ajustam sozinhas.'),
    ('5.', 'Compensação: ficou abaixo no dia 1, a meta do dia 2 sobe; ficou abaixo na semana 1, a meta da semana 2 sobe. Dias passados sem lançamento contam como 0; hoje e dias futuros sem lançamento são considerados cumpridos.'),
    ('6.', 'Feriados nacionais já saem como folga (aba Folgas). Acrescente ali feriados locais ou dias sem expediente.'),
    ('Cores', 'Azul = valor que você digita · Preto = fórmula · Verde = ligação entre abas · Fundo amarelo = célula para preencher.'),
]
for i, (a, b) in enumerate(leg):
    r = 17 + i
    ws.cell(row=r, column=1, value=a).font = f_bold
    ws.cell(row=r, column=1).alignment = Alignment(vertical='top', horizontal='right' if a.endswith('.') else 'left')
    if b:
        c = ws.cell(row=r, column=2, value=b); c.font = f_norm; c.alignment = Alignment(wrap_text=True, vertical='top')
        ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=8)
        ws.row_dimensions[r].height = 30
for col, w in zip('ABCDEFGH', [44, 16, 16, 24, 18, 20, 20, 30]):
    ws.column_dimensions[col].width = w
ws.sheet_view.showGridLines = False

dv = DataValidation(type='list', formula1='"Custo,Número"', allow_blank=False); ws.add_data_validation(dv); dv.add('B10')
dv = DataValidation(type='list', formula1='"Custo,Porcentagem,Número"', allow_blank=False); ws.add_data_validation(dv); dv.add('B11:B13')
dv = DataValidation(type='list', formula1='"Sim,Não"', allow_blank=False); ws.add_data_validation(dv); dv.add('B7')
dv = DataValidation(type='decimal', operator='greaterThanOrEqual', formula1='0', showErrorMessage=True,
                    error='Informe um número maior ou igual a zero.'); ws.add_data_validation(dv); dv.add('B5'); dv.add('C10:C13')
dv = DataValidation(type='date', operator='greaterThan', formula1='36526', showErrorMessage=True,
                    error='Informe uma data válida.'); ws.add_data_validation(dv); dv.add('B4'); dv.add('B6')
ws['B5'].comment = Comment('Dado do usuário: quanto será investido em anúncios no mês.', 'Planilha')
ws['C10'].comment = Comment('Exemplo: R$ 35,00 por lead. Troque pelo custo por lead que você prevê.', 'Planilha')

# ============================================================ Folgas
ws = folgas
ws['A1'] = 'Folgas e feriados'; ws['A1'].font = f_title
ws['A2'] = 'Dias listados aqui não contam como dia útil. Feriados nacionais fixos + Sexta-feira Santa; Carnaval e Corpus Christi ficam de fora (facultativos). Acrescente outros nas linhas livres.'
ws['A2'].font = f_sub
head(ws, 4, ['Data', 'Motivo'])


def easter(y):
    a, b, c = y % 19, y // 100, y % 100
    d, e = b // 4, b % 4
    f = (b + 8) // 25
    g = (b - f + 1) // 3
    h = (19 * a + b - d - g + 15) % 30
    i, k = c // 4, c % 4
    l = (32 + 2 * e + 2 * i - h - k) % 7
    m = (a + 11 * h + 22 * l) // 451
    month = (h + l - 7 * m + 114) // 31
    day = ((h + l - 7 * m + 114) % 31) + 1
    return date(y, month, day)


fixed = {(1, 1): 'Confraternização Universal', (4, 21): 'Tiradentes', (5, 1): 'Dia do Trabalho',
         (9, 7): 'Independência', (10, 12): 'Nossa Sra. Aparecida', (11, 2): 'Finados',
         (11, 15): 'Proclamação da República', (11, 20): 'Consciência Negra', (12, 25): 'Natal'}
hol = []
for y in range(2026, 2031):
    for (m, d), name in fixed.items():
        hol.append((date(y, m, d), name))
    hol.append((easter(y) - timedelta(days=2), 'Sexta-feira Santa'))
hol.sort()
for i, (d, name) in enumerate(hol):
    r = 5 + i
    put(ws, f'A{r}', d, f_input, DATEF, fill_input)
    put(ws, f'B{r}', name, f_input, None, fill_input)
for r in range(5 + len(hol), 5 + len(hol) + 15):   # linhas livres
    put(ws, f'A{r}', None, f_input, DATEF, fill_input)
    put(ws, f'B{r}', None, f_input, None, fill_input)
ws.column_dimensions['A'].width = 16
ws.column_dimensions['B'].width = 34

# ============================================================ Cálculo (motor)
ws = calc
ws['A1'] = 'Cálculo — motor das metas (não precisa mexer)'; ws['A1'].font = f_title
put(ws, 'A2', 'Dias úteis do mês (N)', f_bold); put(ws, 'B2', f'=SUM(C{R0}:C{R1})', f_norm, INT)
put(ws, 'A3', 'Dias úteis já encerrados', f_bold); put(ws, 'B3', f'=COUNTIFS(C{R0}:C{R1},1,A{R0}:A{R1},"<"&{HOJE})', f_norm, INT)
put(ws, 'D2', 'Primeiro dia do mês', f_bold); put(ws, 'E2', f'=DATE(YEAR({B}$B$4),MONTH({B}$B$4),1)', f_link, DATEF)
put(ws, 'D3', 'Segunda-feira da 1ª semana', f_bold); put(ws, 'E3', '=E2-WEEKDAY(E2,2)+1', f_norm, DATEF)
head(ws, 5, ['Data', 'Semana', 'Útil (1/0)', 'Nº do dia útil na semana', 'Dias úteis na semana',
             'Dias úteis do mês a partir da semana', 'Início de semana', 'Semana do calendário (seg-dom)'])
for r in range(R0, R1 + 1):
    i = r - R0
    put(ws, f'A{r}', f'=IF($E$2+{i}<=EOMONTH($E$2,0),$E$2+{i},"")', f_norm, DATEF)
    put(ws, f'H{r}', f'=IF(A{r}="","",INT((A{r}-$E$3)/7)+1)', f_norm, '0')
    put(ws, f'B{r}', f'=IF(A{r}="","",IF(COUNTIFS(H${R0}:H${R1},H{r},C${R0}:C${R1},1)=0,"",IF(H{r}=H{r-1},B{r-1},1+MAX(B${R0-1}:B{r-1}))))', f_norm, '0')
    put(ws, f'C{r}', f'=IF(A{r}="",0,IF(AND(WEEKDAY(A{r},2)<=5,COUNTIF({F}$A$5:$A$300,A{r})=0),1,0))', f_norm, '0')
    put(ws, f'D{r}', f'=IF(C{r}=1,COUNTIFS(B${R0}:B{r},B{r},C${R0}:C{r},1),0)', f_norm, '0')
    put(ws, f'E{r}', f'=IF(OR(A{r}="",B{r}=""),0,COUNTIFS(B${R0}:B${R1},B{r},C${R0}:C${R1},1))', f_norm, '0')
    put(ws, f'F{r}', f'=IF(OR(A{r}="",B{r}=""),0,COUNTIFS(B${R0}:B${R1},">="&B{r},C${R0}:C${R1},1))', f_norm, '0')
    put(ws, f'G{r}', f'=IF(OR(A{r}="",B{r}=""),0,IF(B{r}=B{r-1},0,1))', f_norm, '0')

# métricas: chave, rótulo, é dinheiro, ref do total, coluna do realizado na aba Diária
metrics = [
    ('inv', 'Investimento', True, B + '$B$5', 'G'),
    ('lead', 'Lead', False, B + '$E$10', 'I'),
    ('cot', 'Lead em cotação', False, B + '$E$11', 'K'),
    ('neg', 'Lead em negociação', False, B + '$E$12', 'M'),
    ('ven', 'Venda', False, B + '$E$13', 'O'),
]
blk = {}       # chave -> dict de colunas
col = 10       # coluna J (H = semana do calendário; I em branco)
EPS = '0.000000001'
for key, label, money, tref, rc in metrics:
    if money:
        names = ['P', 'CUM', 'BW', 'W', 'BD', 'MD']
        heads = ['projetado do dia', 'Projetado até o dia anterior', 'Meta-base da semana', 'Meta da semana',
                 'Meta-base do dia', 'Meta do dia']
    else:
        names = ['P', 'CUM', 'AtrW', 'BWb'] + [f'S{k}' for k in range(1, 6)] + [f'B{k}' for k in range(1, 6)] + ['BW', 'W', 'BDr', 'BD', 'MD']
        heads = (['projetado do dia', 'Projetado até o dia anterior', 'Semana atrasada? (1/0)', 'Metas-base das semanas anteriores']
                 + [f'Parte {k} da semana' for k in range(1, 6)] + [f'Parte-base {k} da semana' for k in range(1, 6)]
                 + ['Meta-base da semana', 'Meta da semana', 'Restante base (dias)', 'Meta-base do dia', 'Meta do dia'])
    cols = {n: L(col + j) for j, n in enumerate(names)}
    blk[key] = cols
    c0 = cols['P']
    put(ws, f'{c0}2', f'Total do mês — {label}', f_bold, None, None, left)
    put(ws, f'{c0}3', f'=N({tref})', f_link, MONEY if money else INT)
    ws.merge_cells(f'{c0}2:{cols["MD"]}2')
    head(ws, 5, [(f'{label}: ' if k == 0 else '') + h for k, h in enumerate(heads)], col=col)
    col += len(names)
    T = f'${c0}$3'
    fm = MONEY if money else INT
    for r in range(R0, R1 + 1):
        p = r - 1
        g = lambda n, rr=r: f'{cols[n]}{rr}'
        real = f"{D}{rc}{r}"
        put(ws, g('P'), f'=IF($A{r}="",0,IF({real}<>"",{real},IF(AND($C{r}=1,$A{r}>={HOJE}),{g("MD")},0)))', f_norm, fm)
        put(ws, g('CUM'), 0 if r == R0 else f'={g("CUM", p)}+{g("P", p)}', f_norm, fm)
        done = f'SUMIFS({cols["P"]}${R0-1}:{cols["P"]}{p},$B${R0-1}:$B{p},$B{r})'
        cb = f'SUMIFS({cols["BD"]}${R0-1}:{cols["BD"]}{p},$B${R0-1}:$B{p},$B{r})'
        rem = f'($E{r}-$D{r}+1)'
        if money:
            # contínuo (dinheiro, com centavos)
            put(ws, g('BW'), f'=IF($A{r}="","",IF($G{r}=1,IF($B$2=0,0,{T}*$E{r}/$B$2),{g("BW", p)}))', f_norm, fm)
            need = f'IF($F{r}=0,0,MAX({T}-{g("CUM")},0)*$E{r}/$F{r})'
            put(ws, g('W'), f'=IF($A{r}="","",IF($G{r}=1,IF({COMP},{need},MAX({g("BW")},{need})),{g("W", p)}))', f_norm, fm)
            put(ws, g('BD'), f'=IF($C{r}=1,{g("W")}/$E{r},"")', f_norm, fm)
            nd = f'MAX({g("W")}-{done},0)/{rem}'
            put(ws, g('MD'), f'=IF($C{r}<>1,"",IF({COMP},{nd},MAX({g("BD")},{nd})))', f_norm, fm)
        else:
            # inteiros (pessoas): divisão sequencial, dia a dia (até 5 dias úteis por semana), arredondando;
            # semana atrasada = arredonda para cima. Os passos ficam na 1ª linha da semana.
            first = f'$G{r}=1'
            put(ws, g('BWb'), f'=IF({first},SUMIFS({cols["BW"]}${R0-1}:{cols["BW"]}{p},$G${R0-1}:$G{p},1),"")', f_norm, fm)
            put(ws, g('AtrW'), f'=IF({first},IF({g("CUM")}<{g("BWb")}-{EPS},1,0),"")', f_norm, '0')
            Rneed = f'MAX({T}-{g("CUM")},0)'
            Rbase = f'({T}-{g("BWb")})'
            for k in range(1, 6):
                prev_s = '+'.join(g(f'S{j}') for j in range(1, k)) or '0'
                prev_b = '+'.join(g(f'B{j}') for j in range(1, k)) or '0'
                lk = f'($F{r}-{k - 1})'
                rs = f'({Rneed}-({prev_s}))'
                rb = f'({Rbase}-({prev_b}))'
                put(ws, g(f'S{k}'), f'=IF({first},IF(AND($E{r}>={k},{lk}>0),IF({g("AtrW")}=1,ROUNDUP(ROUND({rs}/{lk},6),0),ROUND({rs}/{lk},0)),0),"")', f_norm, fm)
                put(ws, g(f'B{k}'), f'=IF({first},IF(AND($E{r}>={k},{lk}>0),ROUND({rb}/{lk},0),0),"")', f_norm, fm)
            ssum = '(' + '+'.join(g(f'S{k}') for k in range(1, 6)) + ')'
            bsum = '(' + '+'.join(g(f'B{k}') for k in range(1, 6)) + ')'
            put(ws, g('BW'), f'=IF($A{r}="","",IF({first},{bsum},{g("BW", p)}))', f_norm, fm)
            put(ws, g('W'), f'=IF($A{r}="","",IF({first},IF({COMP},{ssum},MAX({g("BW")},{ssum})),{g("W", p)}))', f_norm, fm)
            put(ws, g('BDr'), f'=IF($A{r}="","",IF({first},{g("BW")},{g("BDr", p)}-N({g("BD", p)})))', f_norm, fm)
            put(ws, g('BD'), f'=IF($C{r}=1,ROUND({g("BDr")}/{rem},0),"")', f_norm, fm)
            xd = f'MAX({g("W")}-{done},0)/{rem}'
            nd = f'IF({done}<{cb}-{EPS},ROUNDUP(ROUND({xd},6),0),ROUND({xd},0))'
            put(ws, g('MD'), f'=IF($C{r}<>1,"",IF({COMP},{nd},MAX({g("BD")},{nd})))', f_norm, fm)
ws.freeze_panes = 'B6'
ws.column_dimensions['A'].width = 24
for c in range(2, col):
    ws.column_dimensions[L(c)].width = 15
ws.sheet_properties.tabColor = '999999'

# ============================================================ Meta Diária
ws = diario
ws['A1'] = 'Meta diária (dias úteis)'; ws['A1'].font = f_title
ws['A2'] = 'Digite só o REALIZADO de cada dia (células amarelas). A meta de cada dia se ajusta sozinha: ficou abaixo, a meta seguinte sobe.'
ws['A2'].font = f_sub
ws['A3'] = 'Verde = bateu a meta do dia · Amarelo = entre 70% e 100% · Vermelho = abaixo de 70% · Custos: verde = dentro do teto, amarelo = até 20% acima, vermelho = acima.'
ws['A3'].font = f_sub
dcols = {}
labels = ['Data', 'Dia', 'Semana', 'Foco do dia', 'Dia útil?']
head(ws, 5, labels)
c = 6
for key, label, money, tref, rc in metrics:
    head(ws, 5, [f'{label}\nMeta do dia', f'{label}\nRealizado'], col=c)
    dcols[key] = (L(c), L(c + 1))
    c += 2
cost_defs = [('Custo por lead\n(real)', 'G', 'I', 10), ('Custo por cotação\n(real)', 'G', 'K', 11),
             ('Custo por negociação\n(real)', 'G', 'M', 12), ('Custo por venda\n(real)', 'G', 'O', 13)]
head(ws, 5, [t[0] for t in cost_defs], col=c)
cost_cols = [L(c + j) for j in range(4)]
put(ws, f'{L(c-1)}4', 'Teto →', f_bold, None, None, Alignment(horizontal='right'))
for j, (t, num, den, brow) in enumerate(cost_defs):
    put(ws, f'{cost_cols[j]}4', f'=IF(ISNUMBER({B}$F${brow}),{B}$F${brow},"")', f_link, MONEY)
ws.row_dimensions[5].height = 42
for r in range(R0, R1 + 1):
    put(ws, f'A{r}', f'=IF({C}A{r}="","",{C}A{r})', f_link, DATEF)
    put(ws, f'B{r}', f'=IF(A{r}="","",CHOOSE(WEEKDAY(A{r},2),"Seg","Ter","Qua","Qui","Sex","Sáb","Dom"))', f_norm, None, None, center)
    put(ws, f'C{r}', f'=IF(A{r}="","",{C}B{r})', f_link, '0', None, center)
    put(ws, f'D{r}', f'=IF(A{r}="","",IF({C}C{r}=1,CHOOSE(WEEKDAY(A{r},2),"Meta da Semana","Foco no Atendimento","Revisão da Semana","Foco no Atendimento","Relatório da Semana"),IF(COUNTIF({F}$A$5:$A$300,A{r})>0,"Feriado / folga","Fim de semana")))', f_norm)
    put(ws, f'E{r}', f'=IF(A{r}="","",IF({C}C{r}=1,"Sim","Não"))', f_norm, None, None, center)
    for key, label, money, tref, rc in metrics:
        mcol, rcol = dcols[key]
        mdcol = blk[key]['MD']
        brow = {'inv': 5, 'lead': 10, 'cot': 11, 'neg': 12, 'ven': 13}[key]
        tcell = f"{B}$B$5" if key == 'inv' else f"{B}$E${brow}"
        put(ws, f'{mcol}{r}', f'=IF(OR($E{r}<>"Sim",NOT(ISNUMBER({tcell}))),"",{C}{mdcol}{r})', f_link, MONEY if money else INT)
        put(ws, f'{rcol}{r}', None, f_input, MONEY if money else INT, fill_input)
    for j, (t, num, den, brow) in enumerate(cost_defs):
        n_, d_ = 'G', cost_defs[j][2]
        put(ws, f'{cost_cols[j]}{r}', f'=IF(AND(ISNUMBER({n_}{r}),ISNUMBER({d_}{r})),IF({d_}{r}>0,{n_}{r}/{d_}{r},""),"")', f_norm, MONEY)
# total
tr = R1 + 1
put(ws, f'A{tr}', 'Total do mês', f_bold)
for key, label, money, tref, rc in metrics:
    mcol, rcol = dcols[key]
    put(ws, f'{rcol}{tr}', f'=IF(COUNT({rcol}{R0}:{rcol}{R1})=0,"",SUM({rcol}{R0}:{rcol}{R1}))', f_bold, MONEY if money else INT)
    put(ws, f'{mcol}{tr}', '', f_bold)
for j in range(4):
    put(ws, f'{cost_cols[j]}{tr}', f'=IF(AND(ISNUMBER(G{tr}),ISNUMBER({cost_defs[j][2]}{tr})),IF({cost_defs[j][2]}{tr}>0,G{tr}/{cost_defs[j][2]}{tr},""),"")', f_bold, MONEY)
# formatação condicional
red = PatternFill('solid', bgColor='F4B6B6', fgColor='F4B6B6')
amb = PatternFill('solid', bgColor='FFE08A', fgColor='FFE08A')
grn = PatternFill('solid', bgColor='B7E4C7', fgColor='B7E4C7')
gry = PatternFill('solid', bgColor='E7E7E7', fgColor='E7E7E7')
last_col = L(c + 3)
ws.conditional_formatting.add(f'A{R0}:{last_col}{R1}', FormulaRule(formula=[f'$E{R0}<>"Sim"'], fill=gry, font=Font(name=FONT, color='888888'), stopIfTrue=True))
for key, label, money, tref, rc in metrics:
    mcol, rcol = dcols[key]
    rng = f'{rcol}{R0}:{rcol}{R1}'
    ws.conditional_formatting.add(rng, FormulaRule(formula=[f'AND(ISNUMBER({rcol}{R0}),ISNUMBER({mcol}{R0}),{rcol}{R0}>={mcol}{R0})'], fill=grn, stopIfTrue=True))
    ws.conditional_formatting.add(rng, FormulaRule(formula=[f'AND(ISNUMBER({rcol}{R0}),ISNUMBER({mcol}{R0}),{rcol}{R0}>={mcol}{R0}*0.7)'], fill=amb, stopIfTrue=True))
    ws.conditional_formatting.add(rng, FormulaRule(formula=[f'AND(ISNUMBER({rcol}{R0}),ISNUMBER({mcol}{R0}))'], fill=red, stopIfTrue=True))
for j in range(4):
    cc = cost_cols[j]
    rng = f'{cc}{R0}:{cc}{R1}'
    ws.conditional_formatting.add(rng, FormulaRule(formula=[f'AND(ISNUMBER({cc}{R0}),ISNUMBER({cc}$4),{cc}{R0}<={cc}$4)'], fill=grn, stopIfTrue=True))
    ws.conditional_formatting.add(rng, FormulaRule(formula=[f'AND(ISNUMBER({cc}{R0}),ISNUMBER({cc}$4),{cc}{R0}<={cc}$4*1.2)'], fill=amb, stopIfTrue=True))
    ws.conditional_formatting.add(rng, FormulaRule(formula=[f'AND(ISNUMBER({cc}{R0}),ISNUMBER({cc}$4))'], fill=red, stopIfTrue=True))
# validações do realizado
dvm = DataValidation(type='decimal', operator='greaterThanOrEqual', formula1='0', showErrorMessage=True,
                     error='Informe um número maior ou igual a zero.')
dvi = DataValidation(type='whole', operator='greaterThanOrEqual', formula1='0', showErrorMessage=True,
                     error='Use um número inteiro (não existe meio lead ou meia venda).')
ws.add_data_validation(dvm); ws.add_data_validation(dvi)
dvm.add(f'G{R0}:G{R1}')
for key in ['lead', 'cot', 'neg', 'ven']:
    dvi.add(f'{dcols[key][1]}{R0}:{dcols[key][1]}{R1}')
ws.freeze_panes = 'F6'
for cc, w in zip('ABCDE', [12, 7, 9, 22, 10]):
    ws.column_dimensions[cc].width = w
for cc in range(6, c + 4):
    ws.column_dimensions[L(cc)].width = 15
ws.sheet_view.showGridLines = False

# ============================================================ Meta Semanal
ws = semanal
ws['A1'] = 'Meta semanal'; ws['A1'].font = f_title
ws['A2'] = 'Semanas de segunda a sexta. A meta de cada semana sobe quando a semana anterior ficou abaixo. O realizado vem da aba Meta Diária.'
ws['A2'].font = f_sub
labels = ['Semana', 'Período', 'Dias úteis']
head(ws, 5, labels)
c = 4
scols = {}
for key, label, money, tref, rc in metrics:
    head(ws, 5, [f'{label}\nMeta da semana', f'{label}\nRealizado'], col=c)
    scols[key] = (L(c), L(c + 1))
    c += 2
head(ws, 5, [t[0] for t in cost_defs], col=c)
scost = [L(c + j) for j in range(4)]
put(ws, f'{L(c-1)}4', 'Teto →', f_bold, None, None, Alignment(horizontal='right'))
for j, (t, num, den, brow) in enumerate(cost_defs):
    put(ws, f'{scost[j]}4', f'=IF(ISNUMBER({B}$F${brow}),{B}$F${brow},"")', f_link, MONEY)
ws.row_dimensions[5].height = 42
for w in range(1, 7):
    r = 5 + w
    ex = f'COUNTIF({C}$B${R0}:$B${R1},$A{r})>0'
    put(ws, f'A{r}', w, f_bold, '0', None, center)
    put(ws, f'B{r}', f'=IF({ex},TEXT(_xlfn.MINIFS({C}$A${R0}:$A${R1},{C}$B${R0}:$B${R1},$A{r},{C}$C${R0}:$C${R1},1),"dd/mm")&" a "&TEXT(_xlfn.MAXIFS({C}$A${R0}:$A${R1},{C}$B${R0}:$B${R1},$A{r},{C}$C${R0}:$C${R1},1),"dd/mm"),"")', f_norm, None, None, center)
    put(ws, f'C{r}', f'=IF({ex},COUNTIFS({C}$B${R0}:$B${R1},$A{r},{C}$C${R0}:$C${R1},1),"")', f_norm, '0', None, center)
    for key, label, money, tref, rc in metrics:
        mcol, rcol = scols[key]
        brow = {'inv': 5, 'lead': 10, 'cot': 11, 'neg': 12, 'ven': 13}[key]
        tcell = f"{B}$B$5" if key == 'inv' else f"{B}$E${brow}"
        wcol = blk[key]['W']
        put(ws, f'{mcol}{r}', f'=IF(AND({ex},ISNUMBER({tcell})),INDEX({C}${wcol}${R0}:${wcol}${R1},MATCH($A{r},{C}$B${R0}:$B${R1},0)),"")', f_link, MONEY if money else INT)
        put(ws, f'{rcol}{r}', f'=IF({ex},IF(COUNTIFS({D}$C${R0}:$C${R1},$A{r},{D}${rc}${R0}:${rc}${R1},"<>")=0,"",SUMIFS({D}${rc}${R0}:${rc}${R1},{D}$C${R0}:$C${R1},$A{r})),"")', f_link, MONEY if money else INT)
    for j, (t, num, den, brow) in enumerate(cost_defs):
        n_ = scols['inv'][1]; d_ = scols[['lead', 'cot', 'neg', 'ven'][j]][1]
        put(ws, f'{scost[j]}{r}', f'=IF(AND(ISNUMBER({n_}{r}),ISNUMBER({d_}{r})),IF({d_}{r}>0,{n_}{r}/{d_}{r},""),"")', f_norm, MONEY)
r = 12
put(ws, f'A{r}', 'Mês', f_bold, None, fill_sec, center)
put(ws, f'B{r}', f'=TEXT({B}$B$4,"mm/yyyy")', f_bold, None, fill_sec, center)
put(ws, f'C{r}', f'={C}B2', f_bold, '0', fill_sec, center)
for key, label, money, tref, rc in metrics:
    mcol, rcol = scols[key]
    brow = {'inv': 5, 'lead': 10, 'cot': 11, 'neg': 12, 'ven': 13}[key]
    tcell = f"{B}$B$5" if key == 'inv' else f"{B}$E${brow}"
    put(ws, f'{mcol}{r}', f'=IF(ISNUMBER({tcell}),{tcell},"")', f_link, MONEY if money else INT, fill_sec)
    put(ws, f'{rcol}{r}', f'=IF(COUNT({rcol}6:{rcol}11)=0,"",SUM({rcol}6:{rcol}11))', f_bold, MONEY if money else INT, fill_sec)
for j in range(4):
    n_ = scols['inv'][1]; d_ = scols[['lead', 'cot', 'neg', 'ven'][j]][1]
    put(ws, f'{scost[j]}{r}', f'=IF(AND(ISNUMBER({n_}{r}),ISNUMBER({d_}{r})),IF({d_}{r}>0,{n_}{r}/{d_}{r},""),"")', f_bold, MONEY, fill_sec)
for key, label, money, tref, rc in metrics:
    mcol, rcol = scols[key]
    rng = f'{rcol}6:{rcol}11'
    ws.conditional_formatting.add(rng, FormulaRule(formula=[f'AND(ISNUMBER({rcol}6),ISNUMBER({mcol}6),{mcol}6>0,{rcol}6>={mcol}6)'], fill=grn, stopIfTrue=True))
    ws.conditional_formatting.add(rng, FormulaRule(formula=[f'AND(ISNUMBER({rcol}6),ISNUMBER({mcol}6),{rcol}6>={mcol}6*0.7)'], fill=amb, stopIfTrue=True))
    ws.conditional_formatting.add(rng, FormulaRule(formula=[f'AND(ISNUMBER({rcol}6),ISNUMBER({mcol}6))'], fill=red, stopIfTrue=True))
for j in range(4):
    cc = scost[j]
    rng = f'{cc}6:{cc}11'
    ws.conditional_formatting.add(rng, FormulaRule(formula=[f'AND(ISNUMBER({cc}6),ISNUMBER({cc}$4),{cc}6<={cc}$4)'], fill=grn, stopIfTrue=True))
    ws.conditional_formatting.add(rng, FormulaRule(formula=[f'AND(ISNUMBER({cc}6),ISNUMBER({cc}$4),{cc}6<={cc}$4*1.2)'], fill=amb, stopIfTrue=True))
    ws.conditional_formatting.add(rng, FormulaRule(formula=[f'AND(ISNUMBER({cc}6),ISNUMBER({cc}$4))'], fill=red, stopIfTrue=True))
ws.freeze_panes = 'D6'
for cc, w in zip('ABC', [10, 16, 11]):
    ws.column_dimensions[cc].width = w
for cc in range(4, c + 4):
    ws.column_dimensions[L(cc)].width = 15
ws.sheet_view.showGridLines = False

# ============================================================ Meta Mensal
ws = mensal
ws['A1'] = 'Meta mensal'; ws['A1'].font = f_title
put(ws, 'A3', 'Mês', f_bold); put(ws, 'B3', f'={B}$B$4', f_link, 'mm/yyyy')
put(ws, 'C3', 'Dias úteis', f_bold); put(ws, 'D3', f'={C}B2', f_link, '0')
put(ws, 'E3', 'Dias encerrados', f_bold); put(ws, 'F3', f'={C}B3', f_link, '0')
put(ws, 'G3', 'Hoje', f_bold); put(ws, 'H3', f'={HOJE}', f_link, DATEF)
head(ws, 5, ['Métrica', 'Meta do mês', 'Realizado', '% da meta', 'Falta', 'Situação'])
rows = [
    ('Investimento (R$)', 'vol', MONEY, f'={B}$B$5', 'G'),
    ('Lead', 'vol', INT, f'={B}$E$10', 'I'),
    ('Custo por lead (teto)', 'cost', MONEY, f'={B}$F$10', ('G', 'I')),
    ('Lead em cotação', 'vol', INT, f'={B}$E$11', 'K'),
    ('Custo por cotação (teto)', 'cost', MONEY, f'={B}$F$11', ('G', 'K')),
    ('Lead em negociação', 'vol', INT, f'={B}$E$12', 'M'),
    ('Custo por negociação (teto)', 'cost', MONEY, f'={B}$F$12', ('G', 'M')),
    ('Venda', 'vol', INT, f'={B}$E$13', 'O'),
    ('Custo por venda (teto)', 'cost', MONEY, f'={B}$F$13', ('G', 'O')),
]
for i, (label, kind, fmt_, meta_f, rc) in enumerate(rows):
    r = 6 + i
    put(ws, f'A{r}', label, f_bold)
    put(ws, f'B{r}', f'=IF(ISNUMBER({meta_f[1:]}),{meta_f[1:]},"")', f_link, fmt_)
    if kind == 'vol':
        put(ws, f'C{r}', f'=IF(COUNT({D}{rc}{R0}:{rc}{R1})=0,"",SUM({D}{rc}{R0}:{rc}{R1}))', f_link, fmt_)
        put(ws, f'D{r}', f'=IF(AND(ISNUMBER(B{r}),B{r}>0,ISNUMBER(C{r})),C{r}/B{r},"")', f_norm, PCT)
        put(ws, f'E{r}', f'=IF(ISNUMBER(B{r}),MAX(B{r}-N(C{r}),0),"")', f_norm, fmt_)
        put(ws, f'F{r}', f'=IF(NOT(ISNUMBER(B{r})),"—",IF($F$3=0,"Aguardando",IF(N(C{r})>=B{r}*$F$3/$D$3,"No ritmo",IF(N(C{r})>=0.7*B{r}*$F$3/$D$3,"Atenção","Abaixo do ritmo"))))', f_bold, None, None, center)
    else:
        n_, d_ = rc
        put(ws, f'C{r}', f'=IF(AND(COUNT({D}{n_}{R0}:{n_}{R1})>0,COUNT({D}{d_}{R0}:{d_}{R1})>0),IF(SUM({D}{d_}{R0}:{d_}{R1})>0,SUM({D}{n_}{R0}:{n_}{R1})/SUM({D}{d_}{R0}:{d_}{R1}),""),"")', f_link, fmt_)
        put(ws, f'D{r}', '', f_norm); put(ws, f'E{r}', '', f_norm)
        put(ws, f'F{r}', f'=IF(NOT(ISNUMBER(B{r})),"—",IF(NOT(ISNUMBER(C{r})),"Aguardando",IF(C{r}<=B{r},"Dentro do teto",IF(C{r}<=B{r}*1.2,"Atenção","Acima do teto"))))', f_bold, None, None, center)
ws.conditional_formatting.add('F6:F14', FormulaRule(formula=['OR(F6="No ritmo",F6="Dentro do teto")'], fill=grn, stopIfTrue=True))
ws.conditional_formatting.add('F6:F14', FormulaRule(formula=['F6="Atenção"'], fill=amb, stopIfTrue=True))
ws.conditional_formatting.add('F6:F14', FormulaRule(formula=['OR(F6="Abaixo do ritmo",F6="Acima do teto")'], fill=red, stopIfTrue=True))
ws['A16'] = ('Meta do mês vem da aba Base de Dados. Custo por unidade (teto) = custo previsto; o realizado é investimento ÷ quantidade lançada. '
             '"No ritmo" compara o realizado com a meta proporcional aos dias úteis já encerrados.')
ws['A16'].font = f_sub
ws.sheet_view.showGridLines = False
for cc, w in zip('ABCDEFGH', [30, 18, 18, 12, 18, 20, 8, 14]):
    ws.column_dimensions[cc].width = w

# lançamentos de exemplo (somente na cópia de teste)
if TEST:
    entries = {  # dia do mês -> (invest, leads, cot, neg, vendas)
        1: (1000, 30, 5, 3, 1), 2: (1200, 20, 4, 2, 0), 5: (1500, 35, 6, 4, 1), 6: (900, 25, 3, 2, 0),
        7: (2000, 50, 9, 6, 2), 8: (800, 20, 4, 2, 1),
    }
    for d_, vals in entries.items():
        r = R0 + d_ - 1
        for (key, label, money, tref, rc), v in zip(metrics, vals):
            diario[f'{rc}{r}'] = v

for w_ in (base, mensal, semanal, diario):
    w_.sheet_properties.tabColor = 'FF6A1A'
from openpyxl.worksheet.properties import PageSetupProperties
for w_ in (base, mensal, semanal, diario):
    w_.page_setup.orientation = 'landscape'
    w_.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=True)
    w_.page_setup.fitToWidth = 1
    w_.page_setup.fitToHeight = 0
wb.active = 0
wb.save(OUT)
print('ok', OUT)
